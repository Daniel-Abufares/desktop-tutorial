"""Shared helpers: config loading, raw file reading, normalisation."""

from __future__ import annotations

import csv
import hashlib
import os
import re
import sys
from datetime import datetime

import yaml

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONFIG_DIR = os.path.join(ROOT, "config")
INBOX_DIR = os.path.join(ROOT, "inbox")
DATA_DIR = os.path.join(ROOT, "data")
REPORTS_DIR = os.path.join(ROOT, "reports")
SAMPLES_DIR = os.path.join(ROOT, "samples")

TRANSACTIONS_CSV = os.path.join(DATA_DIR, "transactions.csv")

FIELDNAMES = [
    "txn_id",
    "date",
    "account",
    "account_label",
    "merchant",
    "merchant_norm",
    "amount",
    "currency",
    "category",
    "notes",
    "source_file",
]

# Encodings Israeli bank/card exports actually show up in, most likely first.
ENCODINGS = ["utf-8-sig", "utf-8", "cp1255", "iso-8859-8", "latin-1"]


def load_yaml(name):
    path = os.path.join(CONFIG_DIR, name)
    with open(path, encoding="utf-8") as fh:
        return yaml.safe_load(fh) or {}


def save_yaml(name, data):
    path = os.path.join(CONFIG_DIR, name)
    with open(path, "w", encoding="utf-8") as fh:
        yaml.safe_dump(data, fh, allow_unicode=True, sort_keys=False)


def die(msg):
    print(f"\n  שגיאה: {msg}\n", file=sys.stderr)
    sys.exit(1)


# --------------------------------------------------------------- reading ----

def read_rows(path):
    """Read any supported statement file into a list of lists of strings."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    if ext in (".csv", ".txt", ".tsv"):
        return _read_csv(path)
    raise ValueError(f"unsupported file type: {ext}")


def _read_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        die("openpyxl חסר. הרץ: pip install -r requirements.txt")
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = []
    for raw in wb[wb.sheetnames[0]].iter_rows(values_only=True):
        rows.append(["" if c is None else str(c).strip() for c in raw])
    wb.close()
    return rows


def _read_csv(path):
    text = None
    for enc in ENCODINGS:
        try:
            with open(path, encoding=enc) as fh:
                text = fh.read()
            break
        except (UnicodeDecodeError, LookupError):
            continue
    if text is None:
        raise ValueError(f"could not decode {path}")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = "\t" if "\t" in sample else ","

    return [[c.strip() for c in row] for row in csv.reader(text.splitlines(), delimiter=delim)]


# ----------------------------------------------------------- normalising ----

_MULTISPACE = re.compile(r"\s+")
# Terminal/branch noise that varies between statements for the same merchant.
_NOISE = re.compile(r"(סניף\s*\d+|מסוף\s*\d+|ח[\"']פ\s*\d+|\bNO\.?\s*\d+\b|\bILS\b|\d{6,})",
                    re.IGNORECASE)


def normalise_merchant(name):
    """Collapse a merchant string so the same shop dedupes across statements."""
    s = str(name or "").strip()
    s = _NOISE.sub(" ", s)
    s = s.replace("*", " ").replace("-", " ").replace("_", " ")
    s = _MULTISPACE.sub(" ", s).strip()
    return s.lower()


def parse_amount(raw):
    """Parse an amount cell into a float. Returns None when there is no number."""
    if raw is None:
        return None
    s = str(raw).strip()
    if not s:
        return None

    negative = s.startswith("(") and s.endswith(")")  # accounting style (250.00)
    s = s.strip("()")

    # Drop currency symbols and letters, keep digits, separators and sign.
    s = re.sub(r"[^\d,.\-+]", "", s)
    if not s or not re.search(r"\d", s):
        return None

    # Israeli exports use "1,234.56". Handle "1.234,56" too.
    if "," in s and "." in s:
        if s.rindex(",") > s.rindex("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        # A single comma is a thousands separator if it has 3 digits after it.
        s = s.replace(",", "") if re.search(r",\d{3}\b", s) else s.replace(",", ".")

    try:
        val = float(s)
    except ValueError:
        return None
    return -val if negative else val


def parse_date(raw, formats):
    """Parse a date cell against the configured formats. Returns ISO or None."""
    if raw is None:
        return None
    s = str(raw).strip()
    if not s:
        return None

    # openpyxl hands back datetimes as "2026-07-12 00:00:00".
    s = s.split(" ")[0]

    for fmt in list(formats) + ["%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y", "%d-%m-%Y"]:
        try:
            dt = datetime.strptime(s, fmt)
            if dt.year < 100:
                dt = dt.replace(year=dt.year + 2000)
            return dt.strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def looks_like_date(cell, formats=()):
    return parse_date(cell, formats) is not None


def txn_id(date, amount, merchant_norm, account):
    key = f"{date}|{amount:.2f}|{merchant_norm}|{account}"
    return hashlib.sha1(key.encode("utf-8")).hexdigest()[:16]


# ------------------------------------------------------------------- io -----

def ensure_dirs():
    for d in (DATA_DIR, REPORTS_DIR):
        os.makedirs(d, exist_ok=True)


def write_transactions(rows, path=None):
    path = path or TRANSACTIONS_CSV
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDNAMES)
        w.writeheader()
        for r in sorted(rows, key=lambda x: (x["date"], x["account"])):
            w.writerow({k: r.get(k, "") for k in FIELDNAMES})


def read_transactions(path=None):
    path = path or TRANSACTIONS_CSV
    if not os.path.exists(path):
        return []
    with open(path, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    for r in rows:
        r["amount"] = float(r["amount"])
    return rows
